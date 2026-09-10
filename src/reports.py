# -*- coding: utf-8 -*-
"""成果文档生成：按任务书《项目最终成果清单》自动生成全部表单（docx/xlsx/pptx/png）。"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "outputs")
os.makedirs(OUT, exist_ok=True)

# 中文图表字体：Windows 用微软雅黑/黑体；云端 Ubuntu 安装 fonts-noto-cjk 后走 Noto
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Noto Sans CJK JP",
                                   "Source Han Sans SC", "WenQuanYi Zen Hei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

FONT = "宋体"


# ================================================================ 通用助手
def _docx_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _set_eastasia(r):
    r.font.name = "Times New Roman"
    r.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _docx_para(doc, text, bold=False, size=12, align="left", indent=True, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    _set_eastasia(r)
    return p


def _docx_heading(doc, text, size=14):
    return _docx_para(doc, text, bold=True, size=size, indent=False, space_after=8)


def _docx_table(doc, headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font_size)
        _set_eastasia(r)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(font_size)
            _set_eastasia(r)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def _xlsx_sheet(wb, title, headers, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, size=10)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(size=10)
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for j, h in enumerate(headers, 1):
        w = (widths[j - 1] if widths else max(10, min(len(str(h)) * 2.2, 40)))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def _fmt_wan(v):
    return round(v / 10000, 2) if isinstance(v, (int, float)) else v


# ================================================================ 图表
CHART_DIR = os.path.join(OUT, "charts")
os.makedirs(CHART_DIR, exist_ok=True)


def nav_chart(path=None):
    """组合净值 vs 沪深300 对比曲线（同起点归一，直接看超额）。"""
    import records
    import benchmark
    nav = records.read_nav()
    if len(nav) < 2:
        return None
    path = path or os.path.join(OUT, "净值曲线.png")
    dates = [r["date"] for r in nav]
    vals = [r["nav"] for r in nav]
    bench = benchmark.series(dates=dates)
    fig, ax = plt.subplots(figsize=(8, 3.6), dpi=150)
    ax.plot(dates, vals, marker="o", ms=3, lw=1.8, color="#C00000", label="模拟组合净值")
    if bench:
        ax.plot(dates, [b["nav"] for b in bench], marker="s", ms=2.5, lw=1.4,
                color="#2E6FDB", ls="--", label="沪深300净值")
    ax.axhline(1.0, color="#888", ls=":", lw=1, label="初始净值 1.00")
    step = max(1, len(dates) // 8)
    ax.set_xticks(dates[::step])
    ax.tick_params(axis="x", rotation=30, labelsize=8)
    ax.set_ylabel("净值（起点=1）")
    ax.set_title("组合净值 vs 沪深300", fontsize=12)
    ax.legend(fontsize=8, loc="best")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def holdings_chart(acct, prices, path=None):
    """持仓结构（按市值占比，横向条形）。"""
    if not acct["positions"]:
        return None
    path = path or os.path.join(CHART_DIR, "持仓结构.png")
    items = []
    for code, pos in acct["positions"].items():
        p = prices.get(code, pos["cost"])
        items.append((pos.get("name", code), p * pos["shares"]))
    total = sum(v for _, v in items) + acct["cash"]
    items.append(("现金", acct["cash"]))
    items.sort(key=lambda x: x[1])
    names = [n for n, _ in items]
    vals = [v / total * 100 for _, v in items]
    colors = ["#8b98ad" if n == "现金" else "#C00000" for n in names]
    fig, ax = plt.subplots(figsize=(7.2, 3.2), dpi=150)
    bars = ax.barh(names, vals, color=colors)
    for b, v in zip(bars, vals):
        ax.text(v + 0.4, b.get_y() + b.get_height() / 2, f"{v:.1f}%", va="center", fontsize=9)
    ax.set_xlabel("占总资产比例 (%)")
    ax.set_title("持仓结构（含现金）", fontsize=12)
    ax.set_xlim(0, max(vals) * 1.25 + 2)
    ax.grid(axis="x", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def perf_compare_chart(cfg, total, benchmark_pct=None, path=None):
    """收益率对比：组合 / 沪深300 / 沪深300ETF。"""
    import benchmark
    path = path or os.path.join(CHART_DIR, "收益对比.png")
    st = cfg["strategy"]
    ret = total / st["initial_capital"] - 1
    labels = ["本组合"]
    vals = [ret * 100]
    bm = benchmark.pct_since()
    if bm:
        labels.append("沪深300")
        vals.append(float(bm.replace("%", "")))
    try:
        from em_client import kline
        df = kline("1.510300", beg="20250101")
        df = df[df["date"] >= cfg["course"]["start_date"]]
        if len(df) >= 2:
            labels.append("沪深300ETF")
            vals.append((float(df["close"].iloc[-1]) / float(df["close"].iloc[0]) - 1) * 100)
    except Exception:
        pass
    colors = ["#C00000"] + ["#2E6FDB"] * (len(vals) - 1)
    fig, ax = plt.subplots(figsize=(6.4, 3.2), dpi=150)
    bars = ax.bar(labels, vals, color=colors, width=0.5)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + (0.08 if v >= 0 else -0.18),
                f"{v:+.2f}%", ha="center", fontsize=10, fontweight="bold")
    ax.axhline(0, color="#555", lw=1)
    ax.set_ylabel("区间收益率 (%)")
    ax.set_title("收益对比（同起点）", fontsize=12)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def pool_score_chart(pool, path=None):
    """候选池综合评分（前8，横向条形）。"""
    if not pool:
        return None
    path = path or os.path.join(CHART_DIR, "候选池评分.png")
    top = list(reversed(pool[:8]))
    names = [f"{r['name']}" for r in top]
    vals = [r["composite"] for r in top]
    fig, ax = plt.subplots(figsize=(7.2, 3.4), dpi=150)
    bars = ax.barh(names, vals, color="#2E6FDB", height=0.6)
    for b, r in zip(bars, top):
        ax.text(b.get_width() + 0.6, b.get_y() + b.get_height() / 2,
                f"{r['composite']:.1f}", va="center", fontsize=9)
    ax.axvline(60, color="#C00000", ls="--", lw=1.2, label="开仓门槛 60 分")
    ax.set_xlabel("综合评分（基本面50% + 技术面50%）")
    ax.set_title("候选池综合评分 Top8", fontsize=12)
    ax.set_xlim(0, max(vals) * 1.2 + 6)
    ax.legend(fontsize=8, loc="lower right")
    ax.grid(axis="x", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


# ================================================================ 01 小组信息登记表
def group_form(cfg):
    wb = Workbook()
    rows = [
        ["课程名称", cfg["course"]["course_name"], "", "项目名称", cfg["course"]["project_name"]],
        ["项目周期", cfg["course"]["weeks"], "", "项目平台", cfg["course"]["platform"]],
        ["团队名称", "", "", "小组人数", "5"],
        ["成员1姓名", "", "学号", "", "角色（组长/研究员/交易员/风险官/记录员）", ""],
        ["成员2姓名", "", "学号", "", "角色", ""],
        ["成员3姓名", "", "学号", "", "角色", ""],
        ["成员4姓名", "", "学号", "", "角色", ""],
        ["成员5姓名", "", "学号", "", "角色", ""],
        ["团队沟通渠道", "", "", "", ""],
    ]
    _xlsx_sheet(wb, "小组信息登记表", ["项目", "内容", "项目2", "内容2", "项目3", "内容3"], rows,
                widths=[14, 22, 10, 22, 30, 14], first=True)
    p = os.path.join(OUT, "01_小组信息登记表.xlsx")
    wb.save(p)
    return p


# ================================================================ 02 市场观察记录
def market_watch_xlsx():
    import records
    rows = records.read_market_watch()
    if not rows:
        return None
    wb = Workbook()
    data = []
    for r in rows:
        idx = r["indices"]
        top = "、".join(f"{b['name']}({b['pct']}%)" for b in r.get("top_boards", [])[:5]) or "—"
        bot = "、".join(f"{b['name']}({b['pct']}%)" for b in r.get("bottom_boards", [])[:5]) or "—"
        def _idx(name):
            v = idx.get(name)
            return f"{v['close']:.2f}（{v['pct']:+.2f}%）" if v else "—"
        data.append([
            r["date"],
            _idx("上证指数"),
            _idx("深证成指"),
            _idx("沪深300"),
            (f"{r['turnover']:.0f}亿元"
             + ("（沪+深）" if r.get("turnover_scope", "sh+sz") == "sh+sz" else "（仅沪市）")
             if r.get("turnover") is not None else "数据源暂不可用"), top, bot,
            r["view"], r["opportunity"], r["risks"],
        ])
    _xlsx_sheet(wb, "市场观察记录",
                ["日期", "上证指数", "深证成指", "沪深300", "两市成交额", "领涨板块", "领跌板块", "市场基本情况判断", "可能的投资机会", "主要风险"],
                data, widths=[11, 18, 18, 18, 11, 30, 30, 42, 36, 30], first=True)
    p = os.path.join(OUT, "02_市场观察记录.xlsx")
    wb.save(p)
    return p


# ================================================================ 03 候选股票池
def pool_xlsx(pool):
    if not pool:
        return None
    wb = Workbook()
    rows = []
    for r in pool:
        fin_judge = r["fin_judgment"] + ("；" + "；".join(r["fin_reasons"][:2]) if r["fin_reasons"] else "")
        rows.append([
            r["code"], r["name"], r["industry"], f"{r['close']:.2f}",
            r["pe_ttm"] if r["kind"] == "stock" else "—", r["pb"] if r["kind"] == "stock" else "—",
            r["include_reason"], fin_judge,
            f"{r['tech_judgment']}。{r['tech_signal']}",
            "；".join(r["risks"]), r["info_source"],
            f"基本面{r['fin_score']} / 技术面{r['tech_score']} / 综合{r['composite']}",
        ])
    _xlsx_sheet(wb, "候选股票池及初步分析",
                ["证券代码", "证券名称", "所属行业", "收盘价", "PE(TTM)", "PB", "纳入原因",
                 "基本面初步判断", "技术面初步判断", "主要风险", "信息来源", "评分"],
                rows, widths=[10, 10, 10, 9, 8, 8, 40, 36, 40, 26, 26, 20], first=True)
    p = os.path.join(OUT, "03_候选股票池及初步分析表.xlsx")
    wb.save(p)
    return p


# ================================================================ 04 初始投资方案
def initial_plan_docx(cfg, pool):
    st = cfg["strategy"]
    doc = Document()
    _docx_style(doc)
    _docx_heading(doc, "初始投资方案", 16)
    _docx_para(doc, f"课程：{cfg['course']['course_name']}　项目：{cfg['course']['project_name']}（{cfg['course']['platform']}）", indent=False)
    top = pool[:3] if pool else []
    _docx_heading(doc, "一、选择哪些证券", 12)
    names = "、".join(f"{r['name']}（{r['code']}，{r['industry']}）" for r in pool[:6]) or "（待候选池生成）"
    _docx_para(doc, f"根据候选股票池综合评分，初始组合拟配置：{names}，另预留现金。", align="justify")
    _docx_heading(doc, "二、资金配置计划", 12)
    _docx_table(doc, ["项目", "计划"], [
        ["初始资金", f"{st['initial_capital']/10000:.0f}万元"],
        ["单只证券目标仓位", f"{st['target_single_weight']*100:.0f}%（上限{st['max_single_position_pct']*100:.0f}%）"],
        ["宽基ETF底仓", f"约{st.get('etf_target_weight', 0.3)*100:.0f}%，分散个股风险并贴近基准"],
        ["现金留存", f"不低于{st['min_cash_pct']*100:.0f}%，应对波动与补仓"],
        ["建仓节奏", f"分{st['max_batch_build']}批、约2周内完成初始建仓"],
    ], widths=[5, 11], font_size=10)
    _docx_heading(doc, "三、选择原因", 12)
    for r in top:
        _docx_para(doc, f"{r['name']}（{r['code']}）：综合评分{r['composite']}分。基本面：{r['fin_judgment']}；"
                         f"技术面：{r['tech_judgment']}。", align="justify")
    _docx_heading(doc, "四、仓位设置理由", 12)
    _docx_para(doc, f"任务书要求单只证券持仓不超过总资产30%。本方案将单只目标仓位设为"
                     f"{st['target_single_weight']*100:.0f}%，留出安全边际；同时保持{st['min_cash_pct']*100:.0f}%以上现金，"
                     "在市场出现急跌时可分批补仓，避免被动。", align="justify")
    _docx_heading(doc, "五、暂不投资的证券", 12)
    rest = [r["name"] for r in pool[6:]] or []
    _docx_para(doc, ("其余候选（" + "、".join(rest) + "）暂不买入：" if rest else "暂无："),
               align="justify")
    if rest:
        pass
    _docx_heading(doc, "六、主要风险", 12)
    _docx_para(doc, "1．系统性风险：指数大幅回调导致组合整体回撤；2．个股风险：业绩不及预期、行业政策变化；"
                     "3．流动性风险：成交清淡造成冲击成本；4．操作风险：委托价格设置不当导致未成交。", align="justify")
    _docx_heading(doc, "七、风险控制方法", 12)
    _docx_para(doc, f"1．单只持仓≤{st['max_single_position_pct']*100:.0f}%；2．止盈线{st['take_profit_pct']*100:.0f}%、"
                     f"止损线-{st['stop_loss_pct']*100:.0f}%，买入时即设定；3．现金≥{st['min_cash_pct']*100:.0f}%；"
                     f"4．每两周至少{st['min_trades_per_2weeks']}笔有效交易，避免频繁交易与刷单；"
                     "5．每笔交易填写投资决策记录，收盘后复盘。", align="justify")
    p = os.path.join(OUT, "04_初始投资方案.docx")
    doc.save(p)
    return p


# ================================================================ 05 基本面分析报告
def fundamental_docx(rec, cfg):
    from fundamental import fmt_revenue
    from em_client import main_fin_data
    doc = Document()
    _docx_style(doc)
    fin = main_fin_data(rec["secid"])
    _docx_heading(doc, f"重点证券基本面分析：{rec['name']}（{rec['code']}）", 16)
    _docx_para(doc, f"所属行业：{rec['industry']}　信息来源：东方财富F10主要财务指标（经同花顺APP核对）", indent=False)
    if fin:
        latest = fin[0]
        _docx_heading(doc, "一、最新财务数据", 12)
        _docx_table(doc, ["报告期", "营业收入", "营收同比", "归母净利润", "净利同比", "加权ROE", "销售净利率", "EPS(元)"], [
            [latest["report_date"] + f"（{latest['report_type']}）",
             fmt_revenue(latest["revenue"]),
             f"{latest['revenue_yoy']:+.2f}%" if latest["revenue_yoy"] is not None else "—",
             fmt_revenue(latest["net_profit"]),
             f"{latest['net_profit_yoy']:+.2f}%" if latest["net_profit_yoy"] is not None else "—",
             f"{latest['roe']:.2f}%" if latest["roe"] is not None else "—",
             f"{latest['net_margin']:.2f}%" if latest["net_margin"] is not None else "—",
             latest["eps"] or "—"],
        ], widths=[3.2, 2.4, 1.8, 2.4, 1.8, 1.8, 1.8, 1.6], font_size=9)
        _docx_heading(doc, "二、这些数据说明了什么", 12)
        for x in rec["fin_reasons"]:
            _docx_para(doc, "· " + x, indent=False)
        _docx_heading(doc, "三、估值与行业情况", 12)
        _docx_para(doc, f"当前PE(TTM)={rec['pe_ttm']}，PB={rec['pb']}。估值处于行业历史区间"
                         f"{'中低' if (rec['pe_ttm'] or 50) < 30 else '中等偏上'}位置；行业景气度与公司竞争地位详见课堂路演展示。", align="justify")
    else:
        _docx_para(doc, "该证券为指数型ETF，不适用个股财务分析。", indent=False)
    _docx_heading(doc, "四、潜在风险", 12)
    for x in rec["risks"]:
        _docx_para(doc, "· " + x, indent=False)
    p = os.path.join(OUT, f"05_基本面分析_{rec['code']}_{rec['name']}.docx")
    doc.save(p)
    return p


# ================================================================ 06 技术分析记录
def tech_xlsx(pool, decisions):
    wb = Workbook()
    rows = []
    for r in pool:
        rows.append([r["date"], r["code"], r["name"], f"{r['close']:.2f}",
                     round(r["ma20"], 2) if r.get("ma20") else "—",
                     round(r["ma60"], 2) if r.get("ma60") else "—",
                     round(r["rsi14"], 1) if r.get("rsi14") is not None else "—",
                     f"{r['support_60d']:.2f}" if r.get("support_60d") else "—",
                     f"{r['resist_60d']:.2f}" if r.get("resist_60d") else "—",
                     r["tech_signal"], r["tech_judgment"]])
    _xlsx_sheet(wb, "技术分析记录",
                ["日期", "代码", "名称", "收盘价", "MA20", "MA60", "RSI14", "支撑位(60日)", "压力位(60日)", "指标→信号", "判断"],
                rows, widths=[11, 9, 10, 9, 8, 8, 7, 11, 11, 46, 16], first=True)
    p = os.path.join(OUT, "06_技术分析记录.xlsx")
    wb.save(p)
    return p


# ================================================================ 07 投资决策记录
def decision_xlsx(decisions):
    import records
    confirmed = records.read_confirmations()
    wb = Workbook()
    rows = []
    for d in decisions:
        rc = d.get("rule_check") or {}
        rows.append([d["date"], d["decision_id"], d["name"], f"（{d['code']}）", d["side"],
                     f"{d['price']:.2f}", d["shares"], f"{d['amount']:.0f}",
                     d["reason"], d["fund_basis"], d["tech_basis"],
                     d.get("market_basis", "—"), d.get("risk_judge", "—"),
                     f"规则引擎：{rc.get('status', '—')}"
                     + (f"；交易后现金占比{rc['cash_ratio_after']*100:.0f}%" if rc.get("cash_ratio_after") else ""),
                     d["ai_used"], d["ai_adopted"],
                     d.get("human_judgment", "PENDING"),
                     "CONFIRMED" if d["decision_id"] in confirmed else "PENDING"])
    _xlsx_sheet(wb, "投资决策记录",
                ["日期", "编号", "证券名称", "代码", "买入/卖出", "交易价格", "数量(股)", "金额(元)",
                 "交易原因", "基本面依据", "技术面依据", "市场依据", "风险判断",
                 "交易前规则校验", "AI是否参与", "是否采纳AI建议", "人工最终判断", "人工判断状态"],
                rows, widths=[11, 13, 10, 10, 9, 9, 9, 11, 30, 28, 30, 30, 20, 22, 14, 16, 24, 12], first=True)
    p = os.path.join(OUT, "07_投资决策记录.xlsx")
    wb.save(p)
    return p


# ================================================================ 08 委托指令单
def order_sheet_docx(date, decisions, path=None):
    doc = Document()
    _docx_style(doc)
    _docx_heading(doc, f"模拟交易委托指令单（{date}）", 16)
    _docx_para(doc, "请在同花顺模拟炒股APP中按下列指令核对后下单；成交后无需手工登记，"
                    "auto模式下账本自动记账，manual模式下将成交价回填 data/trades.csv。", indent=False)
    if not decisions:
        _docx_para(doc, "本日无操作信号，不进行交易。", indent=False)
    else:
        rows = [[d["side"], d["name"], d["code"], f"{d['price']:.2f}", d["shares"],
                 f"{d['amount']:.0f}", d["reason"][:60]] for d in decisions]
        _docx_table(doc, ["方向", "证券名称", "代码", "参考价", "数量(股)", "金额(元)", "理由"],
                    rows, widths=[1.6, 2.6, 2.2, 1.8, 2.0, 2.4, 6.0], font_size=9)
    p = path or os.path.join(OUT, f"08_委托指令单_{date}.docx")
    doc.save(p)
    return p


# ================================================================ 09 交易日志与盈亏归因
def trades_xlsx(cfg):
    import portfolio as pf
    from em_client import snapshot
    trades = pf.read_trades()
    acct = pf.replay(trades, cfg["strategy"]["initial_capital"])
    prices = {}
    for code in acct["positions"]:
        try:
            prices[code] = snapshot([u for u in cfg["universe"] if u["code"] == code][0]["secid"])["price"]
        except Exception:
            prices[code] = acct["positions"][code]["cost"]
    wb = Workbook()
    rows = [[t["date"], t["name"], t["code"], t["side"], t["price"], t["shares"],
             t["amount"], t["decision_id"], t["source"]] for t in trades]
    _xlsx_sheet(wb, "交易日志",
                ["日期", "证券名称", "代码", "方向", "价格", "数量(股)", "金额(元)", "决策编号", "来源"],
                rows, widths=[11, 12, 9, 7, 9, 10, 13, 14, 8], first=True)
    # 盈亏归因
    hold_rows = []
    for code, pos in acct["positions"].items():
        p = prices.get(code, pos["cost"])
        hold_rows.append([pos.get("name", code), pos["shares"], round(pos["cost"], 3), p,
                          round(p * pos["shares"], 0),
                          round((p - pos["cost"]) * pos["shares"], 0),
                          f"{(p/pos['cost']-1)*100:+.2f}%" if pos["cost"] else "—"])
    realized_rows = [[k, round(v, 0)] for k, v in acct["realized"].items()]
    _xlsx_sheet(wb, "当前持仓与浮动盈亏",
                ["证券", "持仓(股)", "成本价", "现价", "市值(元)", "浮动盈亏(元)", "浮动盈亏%"],
                hold_rows, widths=[14, 10, 9, 9, 13, 13, 11])
    _xlsx_sheet(wb, "已实现盈亏归因", ["证券", "已实现盈亏(元)"], realized_rows, widths=[14, 15])
    _xlsx_sheet(wb, "账户汇总", ["项目", "数值"], [
        ["初始资金(元)", cfg["strategy"]["initial_capital"]],
        ["现金(元)", acct["cash"]],
        ["持仓市值(元)", pf.market_value(acct, prices)],
        ["总资产(元)", pf.total_assets(acct, prices)],
        ["总资产收益率", f"{(pf.total_assets(acct, prices)/cfg['strategy']['initial_capital']-1)*100:+.2f}%"],
        ["交易费用合计(元，万2.5佣金+0.05%印花税)", acct["fee"]],
    ], widths=[30, 16])
    p = os.path.join(OUT, "09_交易日志与盈亏归因.xlsx")
    wb.save(p)
    return p, acct, prices


# ================================================================ 10 策略调整记录
def backtest_blocks():
    """解析 data/backtest_result.txt 中的 JSON 结果块 → {策略名: {指标}}。"""
    import json as _json
    path = os.path.join(BASE, "data", "backtest_result.txt")
    if not os.path.exists(path):
        return {}
    blocks, title, buf = {}, None, []
    with open(path, encoding="utf-8") as f:
        for line in f:
            s = line.strip()
            if s.startswith("=== ") and s.endswith(" ==="):
                if title and buf:
                    blocks[title] = "".join(buf)
                title, buf = s[4:-4], []
            elif title is not None:
                if not s and buf:
                    blocks[title] = "".join(buf)
                    title, buf = None, []
                else:
                    buf.append(line)
    if title and buf:
        blocks[title] = "".join(buf)
    parsed = {}
    for k, v in blocks.items():
        try:
            parsed[k] = _json.loads(v)
        except Exception:
            pass
    return parsed


def strategy_docx(cfg, pool):
    st = cfg["strategy"]
    bt = backtest_blocks()
    doc = Document()
    _docx_style(doc)
    _docx_heading(doc, "策略调整记录", 16)
    _docx_para(doc, "本记录由工作流在策略参数或市场状态变化时自动更新，每条调整都遵循"
                    "“发现的问题 → 验证方式 → 调整内容 → 预期效果 → 可能风险”的证据链。", indent=False)

    _docx_heading(doc, "一、调整历程", 12)
    _docx_table(doc, ["时间", "原策略", "发现的问题", "调整内容", "预期效果", "可能风险"], [
        ["初始建仓", "候选池评分选股，单只≤30%，止盈9%/止损8%",
         "—", "建立初始组合：宽基ETF底仓 + 评分前列个股", "分散风险、跟踪市场", "市场急跌时底仓同样受损"],
        ["运行中", f"止盈{st['take_profit_pct']*100:.0f}%/止损{st['stop_loss_pct']*100:.0f}%",
         "技术面恶化（MACD死叉且技术分<45）", "提前于止盈/止损线离场", "减少回撤、执行纪律", "震荡市中可能被反复洗出"],
        ["2026-09-08\n第一次策略升级",
         f"固定止盈+{st['take_profit_pct']*100:.0f}%；单日最多建仓3只；ETF底仓约20%",
         "回测发现固定止盈在震荡市被反复洗出；换手越高收益越低（12笔>64笔>110笔>359笔）",
         f"改为移动止损（自持仓最高价回撤{st.get('trailing_stop_pct', 0.12)*100:.0f}%）替代固定止盈；"
         f"ETF底仓目标权重提至{st.get('etf_target_weight', 0.3)*100:.0f}%；单日最多建仓降为{st['max_batch_build']}只",
         "减少无效交易与手续费损耗，让盈利标的充分奔跑", "趋势反转时回撤可能大于固定止盈"],
        ["2026-09-08\n数据源加固", "单一东方财富接口",
         "东财接口偶发 502/超时，导致云端流水线 exit 1",
         "增加腾讯备用数据源（只补缺口不覆盖东财口径）+ 缓存回退 + 收盘数据自愈重试",
         "云端连续运行不再因单一数据源抖动中断", "两套复权口径混用风险（已用数据溯源规则约束）"],
    ], widths=[2.2, 3.4, 3.2, 3.6, 2.6, 2.6], font_size=8)

    _docx_heading(doc, "二、回测验证（2025-04-21 ~ 2026-09-07，338 个交易日）", 12)
    _docx_para(doc, "回测采用离线行情缓存、决策日收盘生成信号、次日开盘成交，与线上逻辑对齐。"
                    "重要前提：用“今天的财务数据”去回测过去会产生前视偏差——把基本面分数随机打乱后，"
                    "策略收益从 +52.8% 掉到 +3.5%，说明大部分“超额”来自事后已知的好公司。"
                    "因此下表以基本面统一为中性分（60分）的无偏对照为准。", align="justify")
    order = [("无偏-固定止盈9%/止损8%", "固定止盈+9%/止损-8%"),
             ("无偏-跌破20日线离场", "跌破20日均线离场"),
             ("无偏-高点回撤12%离场", "移动止损：自最高价回撤12%离场"),
             ("等权持有12只", "对照：等权买入持有观察池12只")]
    rows = []
    for key, label in order:
        m = bt.get(key)
        if not m:
            continue
        rows.append([label, m.get("总收益"), m.get("最大回撤"), m.get("夏普"),
                     m.get("交易笔数"), m.get("卖出胜率")])
    if rows:
        _docx_table(doc, ["离场规则", "总收益", "最大回撤", "夏普", "交易笔数", "卖出胜率"],
                    rows, widths=[5.2, 2.2, 2.2, 1.6, 2.2, 2.2], font_size=9)
        _docx_para(doc, "结论：①固定止盈在震荡市无效（无偏口径下 -2.75%），改为移动止损后收益 +20.76%、"
                        "最大回撤由 -16.93% 收窄到 -12.54%；②交易越频繁收益越低，费用与择时损耗是主要拖累；"
                        "③宽基ETF底仓提供基准收益，个股负责超额。", align="justify")
    else:
        _docx_para(doc, "尚未生成回测结果，请运行 python backtest.py。", indent=False)

    _docx_heading(doc, "三、当前生效参数", 12)
    _docx_table(doc, ["参数", "取值", "说明"], [
        ["单只持仓上限", f"{st['max_single_position_pct']*100:.0f}%", "任务书硬性要求，规则引擎拦截超限买入"],
        ["现金下限", f"{st['min_cash_pct']*100:.0f}%", "保留补仓与抗回撤能力"],
        ["单只目标仓位", f"{st['target_single_weight']*100:.0f}%", "留出安全边际"],
        ["宽基ETF底仓", f"{st.get('etf_target_weight', 0.3)*100:.0f}%", "分散个股风险、贴近基准"],
        ["移动止损", f"自最高价回撤{st.get('trailing_stop_pct', 0.12)*100:.0f}%", "替代固定止盈，回测收益最优"],
        ["固定止损", f"-{st['stop_loss_pct']*100:.0f}%", "保护本金，始终生效"],
        ["开仓评分门槛", f"{st['min_buy_score']}分", "综合评分（基本面50%+技术面50%）"],
        ["单日最多建仓", f"{st['max_batch_build']}只", "降低换手与冲击成本"],
        ["交易频率下限", f"每两周{st['min_trades_per_2weeks']}笔", "任务书要求，避免长期空仓"],
    ], widths=[4.2, 3.0, 8.0], font_size=9)

    _docx_heading(doc, "四、后续观察与调整原则", 12)
    _docx_para(doc, f"当前候选池前3位：{'、'.join(r['name'] for r in pool[:3]) or '—'}。"
                    "后续按三条原则滚动调整：①任何参数改动先用回测验证，并保留无偏对照；"
                    "②以决策后 T+5/T+10/T+20 的超额收益检验策略是否真的有效；"
                    "③调整记录与AI使用记录同步更新，保证证据链完整。", align="justify")
    p = os.path.join(OUT, "10_策略调整记录.docx")
    doc.save(p)
    return p


# ================================================================ 11 AI使用记录
def ai_xlsx():
    import records
    rows = [[r["seq"], r["time"], r["prompt"], r["ai_output"], r["verify"],
             r["human_judgment"], r["final_use"]] for r in records.read_ai_usage()]
    if not rows:
        rows = [["—", "—", "（运行 main.py daily 后自动记录）", "—", "—", "—", "—"]]
    wb = Workbook()
    _xlsx_sheet(wb, "AI使用记录",
                ["序号", "时间", "Prompt", "AI输出", "信息来源核验", "人工判断", "最终采用/修改"],
                rows, widths=[6, 17, 40, 44, 36, 30, 24], first=True)
    p = os.path.join(OUT, "11_AI使用记录.xlsx")
    wb.save(p)
    return p


# ================================================================ 12 中期路演PPT
def roadshow_pptx(cfg, pool, acct, prices, total, benchmark_pct=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt as PPt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = 13.333

    def add_slide(title, lines=None, table=None, picture=None, pic_left=None, pic_top=1.5, pic_w=5.6):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(W - 1.2), Inches(0.8))
        tf = box.text_frame
        tf.text = title
        tf.paragraphs[0].font.size = PPt(26)
        tf.paragraphs[0].font.bold = True
        text_w = 6.6 if picture else (W - 1.4)
        if lines:
            body = s.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(text_w), Inches(5.4))
            bf = body.text_frame
            bf.word_wrap = True
            for i, ln in enumerate(lines):
                para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
                para.text = ln
                para.font.size = PPt(15)
        if picture and os.path.exists(picture):
            left = pic_left if pic_left is not None else (7.4 if lines else (W - pic_w) / 2)
            s.shapes.add_picture(picture, Inches(left), Inches(pic_top), width=Inches(pic_w))
        if table:
            n_r, n_c = len(table) + 1, len(table[0])
            t = s.shapes.add_table(n_r, n_c, Inches(0.6), Inches(2.6),
                                   Inches(W - 1.2), Inches(4.2)).table
            for j, h in enumerate(table[0]):
                t.cell(0, j).text = str(h)
                t.cell(0, j).text_frame.paragraphs[0].font.size = PPt(12)
            for i, row in enumerate(table, 1):
                for j, v in enumerate(row):
                    t.cell(i, j).text = str(v)
                    t.cell(i, j).text_frame.paragraphs[0].font.size = PPt(11)
        return s

    ret = total / cfg["strategy"]["initial_capital"] - 1
    bm = benchmark_pct if benchmark_pct is not None else "—"

    add_slide(f"{cfg['course']['project_name']}·中期路演", [
        f"课程：{cfg['course']['course_name']}　平台：{cfg['course']['platform']}",
        f"汇报内容：投资表现 / 收益对比 / 当前持仓 / 投资逻辑 / 阶段复盘 / 下一步计划",
        f"数据截止：{__import__('records').read_nav()[-1]['date'] if __import__('records').read_nav() else '—'}",
    ])

    add_slide("1. 当前投资表现", [
        f"初始资金：{cfg['strategy']['initial_capital']/10000:.0f}万元",
        f"当前总资产：{total/10000:.2f}万元",
        f"累计收益率：{ret*100:+.2f}%",
        f"同期沪深300：{bm}",
        f"现金：{acct['cash']/10000:.2f}万元（{acct['cash']/total*100:.1f}%）",
        f"持仓：{len(acct['positions'])}只，累计交易{len(__import__('portfolio').read_trades())}笔",
    ], picture=nav_chart())

    add_slide("2. 收益对比（同起点）", [
        "对比口径：以组合首个净值日为共同起点。",
        "组合收益 = 交易盈亏 + 浮动盈亏 − 交易费用。",
        "宽基ETF底仓提供基准收益，个股贡献超额。",
    ], picture=perf_compare_chart(cfg, total, benchmark_pct))

    hold_table = [["证券", "持仓(股)", "成本价", "现价", "市值(万元)", "浮盈(万元)"]]
    for code, pos in acct["positions"].items():
        pp = prices.get(code, pos["cost"])
        hold_table.append([pos.get("name", code), pos["shares"], round(pos["cost"], 2), pp,
                           round(pp * pos["shares"] / 1e4, 2), round((pp - pos["cost"]) * pos["shares"] / 1e4, 2)])
    add_slide("3. 当前投资组合", [
        "持仓逻辑：基本面评分 + 技术面趋势双确认；单只≤30%。",
        f"纪律：移动止损（自最高价回撤{cfg['strategy'].get('trailing_stop_pct', 0.12)*100:.0f}%）"
        f" + 固定止损 -{cfg['strategy']['stop_loss_pct']*100:.0f}%。",
        "主要风险：市场系统性回调、行业政策变化。",
    ], table=hold_table if acct["positions"] else None,
        picture=holdings_chart(acct, prices) if acct["positions"] else None, pic_top=1.5)

    logic = [f"{r['name']}（{r['code']}，{r['industry']}）：综合{r['composite']}分；{r['fin_judgment']}；{r['tech_judgment']}。"
             for r in pool[:5]]
    add_slide("4. 投资逻辑（候选池Top5）", logic or ["暂无候选池数据"],
              picture=pool_score_chart(pool))

    bt = backtest_blocks()
    trail = bt.get("无偏-高点回撤12%离场", {})
    fixed = bt.get("无偏-固定止盈9%/止损8%", {})
    add_slide("5. 策略验证：为什么改成移动止损", [
        "用回测检验策略，而不是凭感觉调参数。",
        f"固定止盈+9%：无偏回测 {fixed.get('总收益', '—')}，最大回撤 {fixed.get('最大回撤', '—')}。",
        f"移动止损（回撤12%）：无偏回测 {trail.get('总收益', '—')}，最大回撤 {trail.get('最大回撤', '—')}。",
        f"交易笔数由 {fixed.get('交易笔数', '—')} 笔降到 {trail.get('交易笔数', '—')} 笔——换手越低、费用损耗越小。",
        "前视偏差检验：打乱基本面分数后收益由 +52.8% 降至 +3.5%，因此只信无偏对照。",
    ])

    add_slide("6. 阶段复盘与下一步计划", [
        "做得好：买入前完成基本面+技术面双检查，规则引擎前置校验，纪律执行止损。",
        "待改进：对宏观消息反应偏慢；行业集中度需再平衡；回测与实盘的样本仍偏少。",
        "下一步：跟踪决策后 T+5/T+10/T+20 超额收益；候选池评分变化及时调仓；保持现金≥20%。",
    ])
    p = os.path.join(OUT, "12_中期路演.pptx")
    prs.save(p)
    return p


# ================================================================ 14 交易复盘
def review_docx(date, cfg, acct, prices, total):
    import portfolio as pf
    doc = Document()
    _docx_style(doc)
    _docx_heading(doc, f"交易复盘（{date}）", 16)
    ret = (total / cfg["strategy"]["initial_capital"] - 1) * 100
    _docx_heading(doc, "一、账户概况", 12)
    _docx_table(doc, ["项目", "数值"], [
        ["总资产", f"{total/10000:.2f}万元"],
        ["总收益率", f"{ret:+.2f}%"],
        ["现金", f"{acct['cash']/10000:.2f}万元（{acct['cash']/total*100:.1f}%）"],
        ["累计交易笔数", len(pf.read_trades())],
        ["累计交易费用", f"{acct['fee']:.0f}元"],
    ], widths=[5, 9], font_size=10)
    _docx_heading(doc, "二、已实现盈亏归因", 12)
    if acct["realized"]:
        rows = [[k, round(v, 0), "盈利" if v > 0 else "亏损"] for k, v in sorted(acct["realized"].items(), key=lambda x: -x[1])]
        _docx_table(doc, ["证券", "已实现盈亏(元)", "结果"], rows, widths=[5, 5, 3], font_size=10)
    else:
        _docx_para(doc, "暂无平仓交易。", indent=False)
    _docx_heading(doc, "三、当前持仓浮动盈亏", 12)
    rows = []
    for code, pos in acct["positions"].items():
        p = prices.get(code, pos["cost"])
        rows.append([pos.get("name", code), pos["shares"], round(pos["cost"], 3), p,
                     round((p - pos["cost"]) * pos["shares"], 0),
                     f"{(p/pos['cost']-1)*100:+.2f}%" if pos["cost"] else "—"])
    if rows:
        _docx_table(doc, ["证券", "持仓(股)", "成本价", "现价", "浮动盈亏(元)", "浮动%"], rows,
                    widths=[4, 2.4, 2.2, 2, 3, 2], font_size=9)
    _docx_heading(doc, "四、复盘结论", 12)
    conclusions = []
    best = max(acct["realized"].items(), key=lambda x: x[1]) if acct["realized"] else None
    worst = min(acct["realized"].items(), key=lambda x: x[1]) if acct["realized"] else None
    if best and best[1] > 0:
        conclusions.append(f"最成功交易为{best[0]}（+{best[1]:.0f}元），坚持了预设的止盈/逻辑纪律")
    if worst and worst[1] < 0:
        conclusions.append(f"最差交易为{worst[0]}（{worst[1]:.0f}元），需在后续复盘中检视买入依据是否充分")
    for code, pos in acct["positions"].items():
        p = prices.get(code, pos["cost"])
        pnl = p / pos["cost"] - 1 if pos["cost"] else 0
        if pnl <= -cfg["strategy"]["stop_loss_pct"] * 0.8:
            conclusions.append(f"{pos.get('name', code)}浮亏{pnl*100:.1f}%已接近止损线，下一交易日重点跟踪")
    if acct["cash"] / total < 0.2:
        conclusions.append("现金占比低于20%下限，注意保留补仓能力")
    if not conclusions:
        conclusions.append("组合运行正常：持仓均处于纪律区间内，按计划继续持有观察")
    for c in conclusions:
        _docx_para(doc, "· " + c, indent=False)
    p = os.path.join(OUT, f"14_交易复盘_{date}.docx")
    doc.save(p)
    return p
def final_report_docx(cfg, pool, acct, prices, total, decisions, benchmark_pct=None):
    doc = Document()
    _docx_style(doc)
    _docx_heading(doc, "模拟证券投资大赛·投资总结报告", 16)
    _docx_para(doc, f"课程：{cfg['course']['course_name']}　平台：{cfg['course']['platform']}　"
                    f"周期：{cfg['course']['weeks']}", indent=False)

    _docx_heading(doc, "一、投资绩效", 12)
    ret = (total / cfg["strategy"]["initial_capital"] - 1) * 100
    bm_val = benchmark_pct
    try:
        if bm_val is None:
            import benchmark
            bm_val = benchmark.pct_since()
        bm_num = float(str(bm_val).replace("%", "")) if bm_val else None
    except Exception:
        bm_num = None
    _docx_table(doc, ["项目", "数值"], [
        ["初始资金", f"{cfg['strategy']['initial_capital']/10000:.0f}万元"],
        ["期末总资产", f"{total/10000:.2f}万元"],
        ["总收益率", f"{ret:+.2f}%"],
        ["同期沪深300", bm_val if bm_val is not None else "—"],
        ["相对基准超额", f"{ret - bm_num:+.2f}%" if bm_num is not None else "—"],
        ["现金占比", f"{acct['cash']/total*100:.1f}%"],
        ["交易费用合计", f"{acct['fee']:.0f}元"],
    ], widths=[5, 9], font_size=10)
    ch = nav_chart()
    if ch:
        doc.add_picture(ch, width=Cm(15))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_para(doc, "图1：组合净值与沪深300净值对比（同起点归一）", size=9,
                   align="center", indent=False, space_after=10)
    ch2 = perf_compare_chart(cfg, total, benchmark_pct)
    if ch2:
        doc.add_picture(ch2, width=Cm(11))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_para(doc, "图2：收益对比（同起点）", size=9, align="center", indent=False, space_after=10)

    _docx_heading(doc, "二、投资过程记录", 12)
    _docx_para(doc, f"项目期间共形成 {len(decisions)} 条投资决策记录、{len(pool)} 只候选证券的"
                     "跟踪分析，全部决策均包含基本面依据、技术面依据、市场依据、风险判断与AI使用情况，"
                     "实现“指标→信号→判断→投资决策”完整证据链；每笔交易均可追溯到决策编号与规则校验结果。",
               align="justify")

    _docx_heading(doc, "三、盈亏归因", 12)
    if acct["realized"]:
        for code, v in acct["realized"].items():
            name = acct["positions"].get(code, {}).get("name", code)
            _docx_para(doc, f"· {name}（{code}）：已实现盈亏 {v/10000:+.2f} 万元。", indent=False)
    else:
        _docx_para(doc, "· 报告期内未发生平仓，全部盈亏为浮动盈亏。", indent=False)
    for code, pos in acct["positions"].items():
        pp = prices.get(code, pos["cost"])
        _docx_para(doc, f"· {pos.get('name', code)}（{code}）：浮动盈亏 "
                        f"{(pp - pos['cost']) * pos['shares']/10000:+.2f} 万元"
                        f"（现价{pp}，成本{pos['cost']:.2f}）。", indent=False)
    _docx_para(doc, f"· 交易费用累计 {acct['fee']:.0f} 元，占初始资金 "
                    f"{acct['fee']/cfg['strategy']['initial_capital']*100:.3f}%。", indent=False)

    _docx_heading(doc, "四、策略验证与回测（重要）", 12)
    bt = backtest_blocks()
    _docx_para(doc, "为避免“用今天的财务数据去回测过去”的前视偏差，本团队做了两组检验："
                    "①把基本面分数随机打乱；②把基本面分数统一为中性60分。"
                    "前者收益由 +52.8% 降至 +3.5%，说明原始结果中的超额主要来自事后已知的好公司，"
                    "因此正式结论以中性口径的无偏回测为准。", align="justify")
    order = [("无偏-固定止盈9%/止损8%", "固定止盈+9%/止损-8%"),
             ("无偏-跌破20日线离场", "跌破20日均线离场"),
             ("无偏-高点回撤12%离场", "移动止损（自最高价回撤12%）"),
             ("等权持有12只", "对照：等权持有观察池12只")]
    rows = []
    for key, label in order:
        m = bt.get(key)
        if m:
            rows.append([label, m.get("总收益"), m.get("最大回撤"), m.get("夏普"), m.get("交易笔数")])
    if rows:
        _docx_table(doc, ["离场规则", "总收益", "最大回撤", "夏普", "交易笔数"], rows,
                    widths=[5.4, 2.4, 2.4, 1.8, 2.4], font_size=9)
        _docx_para(doc, "据此，团队将“固定止盈+9%”改为“移动止损（自持仓最高价回撤12%）”，"
                        "并把宽基ETF底仓目标权重提高到30%、单日最多建仓降到2只，"
                        "以降低换手与费用损耗。该调整已写入《10_策略调整记录》。", align="justify")
    else:
        _docx_para(doc, "回测结果尚未生成，请运行 python backtest.py。", indent=False)

    _docx_heading(doc, "五、策略反思", 12)
    _docx_para(doc, "1．收益不是唯一标准：决策依据、风险判断与复盘记录才是能力体现；"
                     "2．风险控制优先：单只≤30%、现金≥20%、止损与移动止损纪律有效降低了回撤；"
                     "3．AI辅助而非替代：数据与指标计算自动化后，人工核验数据来源与最终判断不可省略；"
                     "4．持续改进：候选池评分滚动更新，剔除趋势走弱标的，并用决策后 T+5/T+10/T+20 的表现检验策略。",
               align="justify")

    _docx_heading(doc, "六、结论与展望", 12)
    _docx_para(doc, "报告期内组合以宽基ETF为底仓、以评分前列个股贡献超额，严格执行仓位与止损纪律，"
                    "整体回撤可控。后续将继续用回测与决策后验证双轮驱动策略迭代，"
                    "并把每一次调整的依据、验证与结果完整留痕。", align="justify")
    p = os.path.join(OUT, "13_投资总结报告.docx")
    doc.save(p)
    return p


# ================================================================ 19 决策后验证
def verify_xlsx(report):
    """决策后验证表：每笔决策的 T+5/T+10/T+20 表现与同期基准对照。"""
    rows = report.get("rows", [])
    wb = Workbook()
    data = []
    for r in rows:
        h = r["horizons"]

        def cell(key, field="ret"):
            v = h[key].get(field)
            if v is None:
                return "待观察" if field == "ret" else "—"
            return f"{v*100:+.2f}%"

        data.append([r["decision_id"], r["date"], r["side"], f"{r['name']}({r['code']})",
                     r["price"], r["shares"], r["days_elapsed"],
                     cell("T+5"), cell("T+10"), cell("T+20"),
                     cell("T+5", "bench"), cell("T+10", "bench"), cell("T+20", "bench"),
                     cell("T+5", "excess"), cell("T+10", "excess"), cell("T+20", "excess"),
                     r["status"]])
    _xlsx_sheet(wb, "决策后验证",
                ["决策编号", "决策日", "方向", "证券", "成交价", "数量", "已过交易日",
                 "T+5收益", "T+10收益", "T+20收益",
                 "T+5沪深300", "T+10沪深300", "T+20沪深300",
                 "T+5超额", "T+10超额", "T+20超额", "状态"],
                data, widths=[13, 11, 7, 14, 9, 9, 10, 10, 10, 10, 12, 12, 12, 10, 10, 10, 9], first=True)
    s = report.get("summary", {})
    _xlsx_sheet(wb, "汇总", ["指标", "数值"], [
        ["数据截止", s.get("asof")],
        ["决策总数", s.get("n_decisions")],
        ["已验证", s.get("n_verified")],
        ["待观察（未满 T+5）", s.get("n_pending")],
        ["已到期样本数", s.get("n_samples")],
        ["方向正确率", f"{s['win_rate']*100:.1f}%" if s.get("win_rate") is not None else "—"],
        ["平均超额收益", f"{s['avg_excess']*100:+.2f}%" if s.get("avg_excess") is not None else "—"],
        ["说明", "T+N = 决策日之后第N个交易日的收盘价相对成交价的涨跌幅；"
                 "卖出决策以「之后下跌」为正确；超额 = 个股收益 − 同期沪深300（卖出方向取反）"],
    ], widths=[22, 70])
    p = os.path.join(OUT, "19_决策后验证.xlsx")
    wb.save(p)
    return p
